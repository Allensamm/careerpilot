"""Local CV text extraction. No network, no upload, no server."""
import argparse
from pathlib import Path

def extract(path: Path) -> str:
    suffix = path.suffix.lower()
    if suffix == '.txt':
        return path.read_text(encoding='utf-8')
    if suffix == '.pdf':
        from pypdf import PdfReader
        return '\n'.join(page.extract_text() or '' for page in PdfReader(str(path)).pages)
    if suffix == '.docx':
        from docx import Document
        doc = Document(str(path))
        return '\n'.join([p.text for p in doc.paragraphs] + [c.text for t in doc.tables for row in t.rows for c in row.cells])
    raise ValueError('Only .pdf, .docx, and .txt supported')

def main():
    parser = argparse.ArgumentParser(description='Offline CV text extraction')
    parser.add_argument('input', type=Path)
    parser.add_argument('--output', type=Path, help='Destination TXT path; defaults to <name>.extracted.txt')
    args = parser.parse_args()
    if not args.input.is_file() or args.input.stat().st_size > 10_000_000:
        parser.error('Input must be a file smaller than 10 MB')
    output = args.output or args.input.with_suffix('.extracted.txt')
    if output.resolve() == args.input.resolve():
        parser.error('Output must differ from input')
    output.write_text(extract(args.input), encoding='utf-8')
    print(f'Extracted text locally to {output}')

if __name__ == '__main__':
    main()
